"""Insère les images d'un dossier à la fin d'un fichier Word (.docx)."""

from __future__ import annotations

import argparse
import sys
import traceback
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


IMAGE_EXTENSIONS = {".jpeg", ".jpg", ".png", ".gif", ".bmp", ".tiff", ".tif", ".webp"}


def list_images(folder: Path) -> list[Path]:
    images = [
        p
        for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS
    ]
    return sorted(images, key=lambda p: p.name.lower())


def insert_images(docx_path: Path, images: list[Path]) -> None:
    doc = Document(str(docx_path))

    # Démarrer après le contenu existant, sur une nouvelle page.
    doc.add_page_break()

    total = len(images)
    for index, image_path in enumerate(images, start=1):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        # Sans width/height : taille native (DPI de l'image).
        run.add_picture(str(image_path))
        print(f"[{index}/{total}] {image_path.name}")

        if index < total:
            doc.add_page_break()

    doc.save(str(docx_path))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Insère toutes les images d'un dossier à la fin d'un .docx "
            "(une image par page, centrées, taille native). Écrase le fichier source."
        ),
    )
    parser.add_argument(
        "docx",
        type=Path,
        help="Fichier Word .docx à modifier (écrasé)",
    )
    parser.add_argument(
        "images",
        type=Path,
        help="Dossier contenant les images à insérer",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)

    docx_path = args.docx.resolve()
    images_dir = args.images.resolve()

    if not docx_path.is_file() or docx_path.suffix.lower() != ".docx":
        print(f"Fichier .docx invalide: {docx_path}", file=sys.stderr)
        return 1

    if not images_dir.is_dir():
        print(f"Dossier d'images invalide: {images_dir}", file=sys.stderr)
        return 1

    images = list_images(images_dir)
    if not images:
        print(f"Aucune image dans le dossier: {images_dir}", file=sys.stderr)
        return 1

    print(f"DOCX : {docx_path}")
    print(f"Images : {images_dir} ({len(images)} fichier(s))")

    try:
        insert_images(docx_path, images)
    except Exception:
        print(f"Erreur pendant l'insertion dans: {docx_path}", file=sys.stderr)
        traceback.print_exc()
        return 1

    print(f"OK: {len(images)} image(s) ajoutée(s), fichier écrasé.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:
        traceback.print_exc()
        raise SystemExit(1)
